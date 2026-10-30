import os, secrets, csv, io, hashlib, hmac, base64
from types import SimpleNamespace
from datetime import datetime, timezone
from pathlib import Path
import httpx
from fastapi import FastAPI, Depends, HTTPException, status, UploadFile, File, Request
from fastapi.responses import HTMLResponse, Response
from pydantic import BaseModel, Field
from sqlalchemy import create_engine, String, Text, Integer, DateTime, LargeBinary, select, inspect
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, sessionmaker
from pypdf import PdfReader
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from langdetect import detect, LangDetectException

app = FastAPI(title="Vagas IA")
SESSION_SECRET = (os.getenv('APP_SECRET') or os.getenv('APP_PASSWORD', '')).encode('utf-8')
SESSION_COOKIE = 'vagas_ia_session'
SESSION_TTL = 60 * 60 * 24 * 7
APP_USER = os.getenv("APP_USER", "")
APP_PASSWORD = os.getenv("APP_PASSWORD", "")
DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql+psycopg://", 1)
elif DATABASE_URL.startswith("postgresql://"):
    DATABASE_URL = DATABASE_URL.replace("postgresql://", "postgresql+psycopg://", 1)
if DATABASE_URL:
    engine = create_engine(DATABASE_URL, pool_pre_ping=True, connect_args={"sslmode":"require"})
else:
    engine = create_engine("sqlite:////tmp/vagas_ia.db", connect_args={"check_same_thread":False})
Session = sessionmaker(bind=engine, expire_on_commit=False)
BASE = Path(__file__).parent

class Base(DeclarativeBase): pass

class User(Base):
    __tablename__ = "users"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    username: Mapped[str] = mapped_column(String(80), unique=True, index=True)
    email: Mapped[str] = mapped_column(String(255), unique=True, index=True)
    password_hash: Mapped[str] = mapped_column(String(300))
    is_active: Mapped[bool] = mapped_column(default=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda:datetime.now(timezone.utc))

class Profile(Base):
    __tablename__ = "profile"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    user_id: Mapped[int | None] = mapped_column(Integer, nullable=True, unique=True, index=True)
    full_name: Mapped[str] = mapped_column(String(120), default="")
    target_roles: Mapped[str] = mapped_column(String(500), default="")
    target_salary: Mapped[str] = mapped_column(String(100), default="Não definido")
    locations: Mapped[str] = mapped_column(String(300), default="Remoto; híbrido; qualquer lugar do Brasil")
    skills: Mapped[str] = mapped_column(Text, default="")
    resume_text: Mapped[str] = mapped_column(Text, default="")
    resume_filename: Mapped[str] = mapped_column(String(255), default="")
    resume_mime: Mapped[str] = mapped_column(String(100), default="application/pdf")
    resume_pdf: Mapped[bytes | None] = mapped_column(LargeBinary, nullable=True)
    email: Mapped[str] = mapped_column(String(255), default="")
    phone: Mapped[str] = mapped_column(String(80), default="")
    city: Mapped[str] = mapped_column(String(120), default="")
    education: Mapped[str] = mapped_column(String(300), default="")
    experience_summary: Mapped[str] = mapped_column(Text, default="")
    linkedin_url: Mapped[str] = mapped_column(String(500), default="")
    portfolio_url: Mapped[str] = mapped_column(String(500), default="")
    contract_type: Mapped[str] = mapped_column(String(80), default="CLT ou PJ")
    english_level: Mapped[str] = mapped_column(String(80), default="Básico")
    application_policy: Mapped[str] = mapped_column(String(100), default="Automáticas quando permitidas e dados completos")
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda:datetime.now(timezone.utc))

class Job(Base):
    __tablename__ = "jobs"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    user_id: Mapped[int | None] = mapped_column(Integer, nullable=True, index=True)
    url: Mapped[str] = mapped_column(String(1500), index=True)
    title: Mapped[str] = mapped_column(String(500), default="")
    company: Mapped[str] = mapped_column(String(300), default="")
    location: Mapped[str] = mapped_column(String(500), default="")
    source: Mapped[str] = mapped_column(String(100), default="")
    description: Mapped[str] = mapped_column(Text, default="")
    salary: Mapped[str] = mapped_column(String(200), default="")
    contract_type: Mapped[str] = mapped_column(String(120), default="")
    score: Mapped[int] = mapped_column(Integer, default=0)
    status: Mapped[str] = mapped_column(String(60), default="Encontrada")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda:datetime.now(timezone.utc))

class SavedAnswer(Base):
    __tablename__ = "saved_answers"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    user_id: Mapped[int | None] = mapped_column(Integer, nullable=True, index=True)
    question_key: Mapped[str] = mapped_column(String(500), index=True)
    question_text: Mapped[str] = mapped_column(String(1000), default="")
    answer: Mapped[str] = mapped_column(Text, default="")
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda:datetime.now(timezone.utc))

Base.metadata.create_all(engine)
# Upgrade older SQLite deployments in place when the ephemeral database survives a redeploy.
if engine.dialect.name == "sqlite":
    existing_jobs = {col["name"] for col in inspect(engine).get_columns("jobs")}
    with engine.begin() as conn:
        conn.exec_driver_sql("UPDATE profile SET target_salary = 'A partir de R$ 3.500' WHERE target_salary = 'R$ 2.500+'")
        conn.exec_driver_sql("UPDATE profile SET locations = 'Remoto; híbrido; qualquer lugar do Brasil' WHERE locations IN ('Remoto; Brasil', 'Remoto; qualquer lugar do Brasil')")
        if "description" not in existing_jobs:
            conn.exec_driver_sql("ALTER TABLE jobs ADD COLUMN description TEXT NOT NULL DEFAULT ''")
        if "salary" not in existing_jobs:
            conn.exec_driver_sql("ALTER TABLE jobs ADD COLUMN salary VARCHAR(200) NOT NULL DEFAULT ''")
        if "contract_type" not in existing_jobs:
            conn.exec_driver_sql("ALTER TABLE jobs ADD COLUMN contract_type VARCHAR(120) NOT NULL DEFAULT ''")
    existing = {col["name"] for col in inspect(engine).get_columns("profile")}
    with engine.begin() as conn:
        if "resume_text" not in existing:
            conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN resume_text TEXT NOT NULL DEFAULT ''")
        if "resume_filename" not in existing:
            conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN resume_filename VARCHAR(255) NOT NULL DEFAULT ''")
        if "resume_mime" not in existing:
            conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN resume_mime VARCHAR(100) NOT NULL DEFAULT 'application/pdf'")
        if "resume_pdf" not in existing:
            conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN resume_pdf BLOB")
        additions = {"email":"VARCHAR(255) NOT NULL DEFAULT ''","phone":"VARCHAR(80) NOT NULL DEFAULT ''","city":"VARCHAR(120) NOT NULL DEFAULT ''","education":"VARCHAR(300) NOT NULL DEFAULT ''","experience_summary":"TEXT NOT NULL DEFAULT ''","linkedin_url":"VARCHAR(500) NOT NULL DEFAULT ''","portfolio_url":"VARCHAR(500) NOT NULL DEFAULT ''","contract_type":"VARCHAR(80) NOT NULL DEFAULT 'CLT ou PJ'","english_level":"VARCHAR(80) NOT NULL DEFAULT 'Básico'","application_policy":"VARCHAR(100) NOT NULL DEFAULT 'Automáticas quando permitidas e dados completos'"}
        for column, definition in additions.items():
            if column not in existing: conn.exec_driver_sql(f"ALTER TABLE profile ADD COLUMN {column} {definition}")
if engine.dialect.name == "postgresql":
    with engine.begin() as conn:
        conn.exec_driver_sql("UPDATE profile SET target_salary = 'A partir de R$ 3.500' WHERE target_salary = 'R$ 2.500+'")
        conn.exec_driver_sql("UPDATE profile SET locations = 'Remoto; híbrido; qualquer lugar do Brasil' WHERE locations IN ('Remoto; Brasil', 'Remoto; qualquer lugar do Brasil')")
        conn.exec_driver_sql("ALTER TABLE jobs ADD COLUMN IF NOT EXISTS salary VARCHAR(200) NOT NULL DEFAULT ''")
        conn.exec_driver_sql("ALTER TABLE jobs ADD COLUMN IF NOT EXISTS contract_type VARCHAR(120) NOT NULL DEFAULT ''")
        conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN IF NOT EXISTS resume_pdf BYTEA")
        conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN IF NOT EXISTS resume_mime VARCHAR(100) NOT NULL DEFAULT 'application/pdf'")
        conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN IF NOT EXISTS email VARCHAR(255) NOT NULL DEFAULT ''")
        conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN IF NOT EXISTS phone VARCHAR(80) NOT NULL DEFAULT ''")
        conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN IF NOT EXISTS city VARCHAR(120) NOT NULL DEFAULT ''")
        conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN IF NOT EXISTS education VARCHAR(300) NOT NULL DEFAULT ''")
        conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN IF NOT EXISTS experience_summary TEXT NOT NULL DEFAULT ''")
        conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN IF NOT EXISTS linkedin_url VARCHAR(500) NOT NULL DEFAULT ''")
        conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN IF NOT EXISTS portfolio_url VARCHAR(500) NOT NULL DEFAULT ''")
        conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN IF NOT EXISTS contract_type VARCHAR(80) NOT NULL DEFAULT 'CLT ou PJ'")
        conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN IF NOT EXISTS english_level VARCHAR(80) NOT NULL DEFAULT 'Básico'")
        conn.exec_driver_sql("ALTER TABLE profile ADD COLUMN IF NOT EXISTS application_policy VARCHAR(100) NOT NULL DEFAULT 'Automáticas quando permitidas e dados completos'")

def hash_password(password: str) -> str:
    salt=secrets.token_bytes(16)
    digest=hashlib.scrypt(password.encode("utf-8"),salt=salt,n=2**14,r=8,p=1)
    return "scrypt$"+base64.urlsafe_b64encode(salt).decode()+"$"+base64.urlsafe_b64encode(digest).decode()

def verify_password(password: str, encoded: str) -> bool:
    try:
        algorithm,salt_b64,digest_b64=encoded.split("$",2)
        if algorithm!="scrypt": return False
        salt=base64.urlsafe_b64decode(salt_b64.encode())
        expected=base64.urlsafe_b64decode(digest_b64.encode())
        actual=hashlib.scrypt(password.encode("utf-8"),salt=salt,n=2**14,r=8,p=1)
        return hmac.compare_digest(actual,expected)
    except Exception:
        return False

# Add tenant ownership columns and migrate all existing records to the legacy owner.
with engine.begin() as conn:
    for table in ("profile","jobs","saved_answers"):
        cols={c["name"] for c in inspect(engine).get_columns(table)}
        if "user_id" not in cols:
            conn.exec_driver_sql(f"ALTER TABLE {table} ADD COLUMN user_id INTEGER")
    # Old indexes enforced global uniqueness. The same job/question must be saveable by separate users.
    conn.exec_driver_sql("DROP INDEX IF EXISTS ix_jobs_url")
    conn.exec_driver_sql("DROP INDEX IF EXISTS ix_saved_answers_question_key")

legacy_user_id=None
if APP_USER and APP_PASSWORD:
    with Session() as db:
        legacy=db.scalar(select(User).where(User.username==APP_USER))
        if legacy is None:
            legacy=User(username=APP_USER,email=(APP_USER.lower() if "@" in APP_USER else APP_USER.lower()+"@legacy.local"),password_hash=hash_password(APP_PASSWORD))
            db.add(legacy)
            db.flush()
        legacy_user_id=legacy.id
        db.commit()

with engine.begin() as conn:
    if legacy_user_id is not None:
        for table in ("profile","jobs","saved_answers"):
            conn.exec_driver_sql(f"UPDATE {table} SET user_id = {int(legacy_user_id)} WHERE user_id IS NULL")
    conn.exec_driver_sql("CREATE UNIQUE INDEX IF NOT EXISTS ux_profile_user_id ON profile (user_id)")
    conn.exec_driver_sql("CREATE INDEX IF NOT EXISTS ix_jobs_user_id ON jobs (user_id)")
    conn.exec_driver_sql("CREATE INDEX IF NOT EXISTS ix_jobs_user_url ON jobs (user_id, url)")
    conn.exec_driver_sql("CREATE INDEX IF NOT EXISTS ix_saved_answers_user_id ON saved_answers (user_id)")
    conn.exec_driver_sql("CREATE INDEX IF NOT EXISTS ix_saved_answers_user_question ON saved_answers (user_id, question_key)")

with Session() as db:
    unassigned=0
    for model in (Profile,Job,SavedAnswer):
        unassigned += db.query(model).filter(model.user_id.is_(None)).count()
    if unassigned:
        raise RuntimeError("Existem dados antigos sem proprietário. Configure APP_USER e APP_PASSWORD no Render para migrar a conta atual.")
    if legacy_user_id is not None and db.scalar(select(Profile).where(Profile.user_id==legacy_user_id)) is None:
        db.add(Profile(user_id=legacy_user_id))
        db.commit()

def session_token(user_id: int) -> str:
    expires=int(datetime.now(timezone.utc).timestamp())+SESSION_TTL
    payload=f"{user_id}.{expires}"
    sig=hmac.new(SESSION_SECRET,payload.encode(),hashlib.sha256).hexdigest()
    return payload+"."+sig

def session_user_id(token: str | None):
    if not token or not SESSION_SECRET: return None
    try:
        uid,expiry,sig=token.split(".",2)
        payload=uid+"."+expiry
        expected=hmac.new(SESSION_SECRET,payload.encode(),hashlib.sha256).hexdigest()
        if not hmac.compare_digest(sig,expected) or int(expiry)<int(datetime.now(timezone.utc).timestamp()): return None
        return int(uid)
    except Exception:
        return None

def require_auth(request: Request):
    user_id=session_user_id(request.cookies.get("vagas_ia_session"))
    if user_id is None:
        raise HTTPException(status_code=401,detail="Faça login para continuar.")
    with Session() as db:
        user=db.get(User,user_id)
        if user is None or not user.is_active:
            raise HTTPException(status_code=401,detail="Sessão inválida. Entre novamente.")
        return SimpleNamespace(id=user.id,username=user.username,email=user.email)

def profile_dict(p):
    return {"full_name":p.full_name,"target_roles":p.target_roles,"target_salary":p.target_salary,"locations":p.locations,"skills":p.skills,"email":p.email,"phone":p.phone,"city":p.city,"education":p.education,"experience_summary":p.experience_summary,"linkedin_url":p.linkedin_url,"portfolio_url":p.portfolio_url,"contract_type":p.contract_type,"english_level":p.english_level,"application_policy":p.application_policy,"resume_filename":p.resume_filename,"resume_text_length":len(p.resume_text or ""),"updated_at":p.updated_at.isoformat() if p.updated_at else None}

EXCLUDE = ["senior data analyst","data analyst","data scientist","software engineer","senior developer","principal engineer","account executive","financial analyst","accountant","nurse","medical doctor"]

def is_portuguese_listing(title, description):
    """Strictly accept jobs whose description language is Portuguese."""
    sample = " ".join((description or "").split())
    if len(sample) < 80:
        return False
    try:
        return detect(sample[:12000]) == "pt"
    except LangDetectException:
        return False



def minimum_salary_brl(value):
    import re
    raw=(value or "").lower()
    m=re.search(r"(?:r\$|brl)\s*([\d.]+(?:,[\d]{1,2})?)", raw)
    if not m: return None
    try: return float(m.group(1).replace(".","").replace(",","."))
    except ValueError: return None
def salary_is_below_minimum(salary, minimum=None):
    """Filtra apenas faixa mensal em BRL cujo teto está abaixo do mínimo."""
    import re
    raw=(salary or "").lower()
    if minimum is None or not raw or not ("r$" in raw or "brl" in raw): return False
    if any(x in raw for x in ["/ano", "anual", "year", "/hour", "/hora", "hourly", "por hora"]): return False
    values=[]
    for item in re.findall(r"(?:r\$|brl)\s*([\d.]+(?:,[\d]{1,2})?)", raw):
        try: values.append(float(item.replace(".","").replace(",",".")))
        except ValueError: pass
    return bool(values) and max(values) < minimum

def remote_location_allowed(location, description=""):
    """Aceita remoto/híbrido; descarta apenas modalidade presencial explícita e restrições incompatíveis."""
    loc=(location or "").lower()
    if any(x in loc for x in ["on-site", "onsite", "on site", "presencial", "in-office", "in office", "no escritório", "no escritorio"]): return False
    text=loc+" "+(description or "")[:1800].lower()
    restricted=["united states only", "usa only", "uk only", "canada only", "europe only", "must be located in the united states", "only in the united states", "somente nos estados unidos"]
    return not any(x in text for x in restricted)

def location_verification_label(location):
    loc=(location or "").lower()
    terms=["brazil", "brasil", "worldwide", "anywhere", "global", "latin america", "south america", "all countries", "qualquer lugar"]
    return "Elegibilidade Brasil indicada pela fonte" if any(x in loc for x in terms) else "Confirmar se aceita profissionais no Brasil"

def score_job(title, location, roles="", skills="", resume_text="", description=""):
    t=(title or "").lower(); loc=(location or "").lower()
    desc=(description or "").lower()
    if not is_portuguese_listing(title, description): return 0
    if not remote_location_allowed(location, description): return 0
    text=(t+" "+desc)
    # Exclude clear Spanish-bilingual roles, not every job mentioning Spanish in unrelated context.
    spanish_markers=["spanish","español","espanhol","língua espanhola","lingua espanhola"]
    bilingual_markers=["bilingual","bilingüe","bilíngue","fluent","fluency","native","advanced","required","proficiency","speaking","fluente","fluência","fluencia"]
    if any(x in t for x in ["bilingual (spanish)","bilingual spanish","spanish-speaking","spanish speaking","customer support specialist, bilingual"]): return 0
    if any(x in text for x in spanish_markers) and any(x in text for x in bilingual_markers): return 0
    english_required=["fluent in english","english fluency","fluency in english","native-level english","advanced english required","english is required","english required","professional english proficiency","inglês fluente","ingles fluente","fluência em inglês","fluencia en inglés"]
    if any(term in text for term in english_required): return 0
    if any(x in t for x in EXCLUDE): return 0
    role_tokens=[x.strip().lower() for x in (roles or "").split(";") if x.strip()]
    skill_tokens=[x.strip().lower() for x in (skills or "").split(";") if len(x.strip())>=4]
    relevant_terms=["marketplace","e-commerce","ecommerce","social media","customer support","customer service","virtual assistant","content","marketing assistant","assistente","atendimento","suporte","catalog","product listing","community manager","digital marketing","marketing coordinator","administrative assistant","customer experience","product catalog","e-commerce assistant","ecommerce assistant","marketplace analyst","marketplace assistant","seller support","operations assistant","sales support","customer success","support specialist","content creator","social media assistant","marketing assistant","product data","listing specialist"]
    cv=(resume_text or "").lower()
    cv_terms=["shopee","mercado livre","tiktok shop","shein","temu","amazon","marketplace","e-commerce","ecommerce","social media","instagram","tiktok","canva","excel","catalog","product listing","ads","anúncios","atendimento","customer support","customer service","content","marketing"]
    role_match=any(x in text for x in role_tokens)
    skill_match=any(x in text for x in skill_tokens)
    relevant_match=any(x in text for x in relevant_terms)
    cv_match=any(x in cv and x in text for x in cv_terms)
    if not (role_match or skill_match or relevant_match or cv_match): return 0
    score=15
    for token in role_tokens:
        if token in text: score=min(100,score+20)
    for token in skill_tokens:
        if token in text: score=min(100,score+8)
    for token in relevant_terms:
        if token in text: score=min(100,score+8)
    for token in cv_terms:
        if token in cv and token in text: score=min(100,score+5)
    if any(x in loc for x in ["brazil","brasil","latin america","south america","worldwide","anywhere","global","all countries"]): score=min(100,score+10)
    return min(100,score)
class StatusUpdate(BaseModel):
    status: str
class RegisterPayload(BaseModel):
    username: str = Field(min_length=3,max_length=80,pattern=r"^[a-zA-Z0-9_.-]+$")
    email: str = Field(min_length=5,max_length=255)
    password: str = Field(min_length=10,max_length=128)

class LoginPayload(BaseModel):
    identifier: str = Field(min_length=1,max_length=255)
    password: str = Field(min_length=1,max_length=128)

class ProfileUpdate(BaseModel):
    full_name: str = Field(default="", max_length=120)
    target_roles: str = Field(default="", max_length=500)
    target_salary: str = Field(default="Não definido", max_length=100)
    locations: str = Field(default="Remoto; híbrido; qualquer lugar do Brasil", max_length=300)
    skills: str = Field(default="", max_length=4000)
    email: str = Field(default="", max_length=255)
    phone: str = Field(default="", max_length=80)
    city: str = Field(default="", max_length=120)
    education: str = Field(default="", max_length=300)
    experience_summary: str = Field(default="", max_length=6000)
    linkedin_url: str = Field(default="", max_length=500)
    portfolio_url: str = Field(default="", max_length=500)
    contract_type: str = Field(default="CLT ou PJ", max_length=80)
    english_level: str = Field(default="Básico", max_length=80)
    application_policy: str = Field(default="Automáticas quando permitidas e dados completos", max_length=100)

class AnswerPayload(BaseModel):
    question: str = Field(min_length=3, max_length=1000)
    answer: str = Field(min_length=1, max_length=6000)

@app.get("/api/onboarding/status")
def onboarding_status(user=Depends(require_auth)):
    with Session() as db:
        p=db.scalar(select(Profile).where(Profile.user_id==user.id))
        checks=[("Nome completo",bool((p.full_name or "").strip())),("E-mail",bool((p.email or "").strip())),("Telefone",bool((p.phone or "").strip())),("Cidade/UF",bool((p.city or "").strip())),("Formação",bool((p.education or "").strip())),("Currículo em PDF",bool((p.resume_text or "").strip())),("Experiências profissionais",bool((p.experience_summary or "").strip()))]
        missing=[label for label,done in checks if not done]
        return {"ready":len(missing)==0,"completed":len(checks)-len(missing),"total":len(checks),"missing":missing,"checks":[{"label":label,"complete":done} for label,done in checks],"message":"Perfil completo para preparar candidaturas; cada portal ainda pode exigir informações adicionais." if not missing else "Complete os dados abaixo antes de preparar candidaturas."}

@app.get("/api/answers")
def list_saved_answers(user=Depends(require_auth)):
    with Session() as db:
        answers=db.scalars(select(SavedAnswer).where(SavedAnswer.user_id==user.id).order_by(SavedAnswer.question_text.asc())).all()
        return [{"id":a.id,"question":a.question_text,"answer":a.answer,"updated_at":a.updated_at.isoformat() if a.updated_at else None} for a in answers]

@app.post("/api/answers")
def save_answer(payload: AnswerPayload, user=Depends(require_auth)):
    question=" ".join(payload.question.split()).strip()
    answer=payload.answer.strip()
    if not answer: raise HTTPException(400,"A resposta não pode ficar vazia.")
    key=" ".join(question.casefold().split())
    with Session() as db:
        saved=db.scalar(select(SavedAnswer).where(SavedAnswer.user_id==user.id,SavedAnswer.question_key==key))
        if saved:
            saved.question_text=question; saved.answer=answer; saved.updated_at=datetime.now(timezone.utc)
        else:
            saved=SavedAnswer(user_id=user.id,question_key=key,question_text=question,answer=answer,updated_at=datetime.now(timezone.utc)); db.add(saved)
        db.commit()
        return {"ok":True,"message":"Resposta salva para reutilização quando a pergunta for equivalente.","question":saved.question_text}

@app.get("/", response_class=HTMLResponse)
def home():
    return (BASE / "static" / "login.html").read_text(encoding="utf-8")

@app.get("/app", response_class=HTMLResponse)
def dashboard(user=Depends(require_auth)):
    return (BASE / "static" / "index.html").read_text(encoding="utf-8")

@app.post("/api/auth/register")
def register(payload: RegisterPayload, request: Request, response: Response):
    if not SESSION_SECRET: raise HTTPException(503,"Configure APP_SECRET no Render.")
    username=payload.username.strip().lower()
    email=payload.email.strip().lower()
    if "@" not in email or "." not in email.rsplit("@",1)[-1]:
        raise HTTPException(400,"Informe um e-mail válido.")
    with Session() as db:
        if db.scalar(select(User).where((User.username==username)|(User.email==email))):
            raise HTTPException(409,"Usuário ou e-mail já cadastrado.")
        user=User(username=username,email=email,password_hash=hash_password(payload.password))
        db.add(user); db.flush()
        db.add(Profile(user_id=user.id))
        db.commit()
        uid=user.id
    response.set_cookie("vagas_ia_session",session_token(uid),max_age=SESSION_TTL,httponly=True,secure=request.url.scheme=="https",samesite="lax",path="/")
    return {"ok":True,"username":username}

@app.post("/api/auth/login")
def login(payload: LoginPayload, request: Request, response: Response):
    identifier=payload.identifier.strip().lower()
    with Session() as db:
        user=db.scalar(select(User).where((User.email==identifier)|(User.username==identifier)))
        if user is None or not user.is_active or not verify_password(payload.password,user.password_hash):
            raise HTTPException(401,"Usuário ou senha incorretos.")
        uid=user.id; username=user.username
    response.set_cookie("vagas_ia_session",session_token(uid),max_age=SESSION_TTL,httponly=True,secure=request.url.scheme=="https",samesite="lax",path="/")
    return {"ok":True,"username":username}

@app.get("/api/auth/me")
def auth_me(user=Depends(require_auth)):
    return {"authenticated":True,"username":user.username}

@app.post("/api/auth/logout")
def logout(response: Response):
    response.delete_cookie("vagas_ia_session",path="/",httponly=True,samesite="lax")
    return {"ok":True}

@app.get("/api/health")
def health():
    return {"ok":True,"service":"vagas-ia-agent","persistent_database":bool(DATABASE_URL)}

@app.get("/api/profile")
def get_profile(user=Depends(require_auth)):
    with Session() as db: return profile_dict(db.scalar(select(Profile).where(Profile.user_id==user.id)))

@app.put("/api/profile")
def save_profile(payload: ProfileUpdate, user=Depends(require_auth)):
    with Session() as db:
        p=db.scalar(select(Profile).where(Profile.user_id==user.id))
        if p is None: raise HTTPException(404,"Perfil não encontrado")
        for key,value in payload.model_dump().items(): setattr(p,key,value.strip())
        p.updated_at=datetime.now(timezone.utc); db.commit()
        return {"ok":True,"message":"Perfil salvo"}

@app.post("/api/profile/resume")
async def upload_resume(file: UploadFile = File(...), user=Depends(require_auth)):
    if not file.filename or not file.filename.lower().endswith(".pdf"):
        raise HTTPException(400,"Envie seu currículo em PDF.")
    data=await file.read()
    if len(data)>8*1024*1024: raise HTTPException(413,"O PDF deve ter até 8 MB.")
    try:
        reader=PdfReader(io.BytesIO(data))
        text="\n".join(page.extract_text() or "" for page in reader.pages).strip()
    except Exception: raise HTTPException(400,"Não consegui ler esse PDF. Tente exportar o currículo novamente em PDF.")
    if not text: raise HTTPException(400,"O PDF parece ser uma imagem sem texto selecionável. Envie um PDF com texto.")
    with Session() as db:
        p=db.scalar(select(Profile).where(Profile.user_id==user.id))
        if p is None: raise HTTPException(404,"Perfil não encontrado")
        p.resume_text=text[:60000]; p.resume_filename=file.filename[:255]; p.resume_mime="application/pdf"; p.resume_pdf=data; p.updated_at=datetime.now(timezone.utc); db.commit()
    return {"ok":True,"filename":file.filename,"characters_extracted":len(text),"message":"Currículo lido e salvo no banco configurado."}


class TailorRequest(BaseModel):
    job_id: int
    format: str = Field(default="docx", pattern="^(docx|txt)$")

def resume_sections(text):
    lines=[line.strip() for line in (text or "").replace("\\r","").splitlines() if line.strip()]
    headings={"experiência profissional","experiencia profissional","experiência","experiencia","formação acadêmica","formacao academica","educação","educacao","competências","competencias","habilidades","cursos","idiomas","resumo profissional","objetivo"}
    sections=[]; current="Resumo do currículo"; buf=[]
    for line in lines:
        norm=line.lower().strip(" :")
        if norm in headings:
            if buf: sections.append((current,buf)); buf=[]
            current=line
        else: buf.append(line)
    if buf: sections.append((current,buf))
    return sections

def tailor_resume_text(profile, job):
    original=(profile.resume_text or "").strip()
    if len(original)<100:
        raise HTTPException(400,"Envie primeiro seu currículo em PDF com texto selecionável.")
    title=job.title or "Vaga"
    description=(job.description or "")
    tokens=["marketplace","e-commerce","ecommerce","social media","instagram","tiktok","shopee","mercado livre","catalog","catálogo","anúncios","ads","atendimento","customer support","excel","canva","conteúdo","content","logística","estoque","produto","product listing","gestão","gestao"]
    original_lower=original.lower()
    desc_lower=description.lower()
    title_lower=title.lower()
    matched=[x for x in tokens if x in original_lower and (x in desc_lower or x in title_lower)]
    # Rearrange only existing, verifiable text. Do not fabricate achievements, skills, education, or languages.
    sections=resume_sections(original)
    ordered=[]
    for name,lines in sections:
        joined=" ".join(lines).lower()
        points=sum(1 for x in matched if x in joined)
        ordered.append((points,name,lines))
    ordered.sort(key=lambda x:x[0],reverse=True)
    out=[profile.full_name or "Currículo profissional",f"Versão direcionada: {title}",f"Empresa: {job.company or 'Não informada'}","", "Resumo de direcionamento",f"Esta versão reorganiza o conteúdo existente para destacar termos presentes no anúncio: {', '.join(matched[:12]) if matched else 'nenhum termo em comum identificado'}.",""]
    for _,name,lines in ordered:
        out.append(name)
        out.extend(lines)
        out.append("")
    out += ["Nota de integridade","Documento reorganizado a partir do currículo original. Revise o conteúdo antes de enviar; nenhuma experiência ou qualificação nova foi criada."]
    return "\n".join(out)

@app.post("/api/jobs/{job_id}/tailor")
def tailor_resume(job_id:int, user=Depends(require_auth), format:str="docx"):
    if format not in {"docx","pdf","txt"}: raise HTTPException(400,"Formato deve ser docx, pdf ou txt.")
    with Session() as db:
        job=db.scalar(select(Job).where(Job.id==job_id,Job.user_id==user.id)); profile=db.scalar(select(Profile).where(Profile.user_id==user.id))
        if not job or not profile: raise HTTPException(404,"Vaga não encontrada")
        text=tailor_resume_text(profile,job)
    filename=f"curriculo_personalizado_vaga_{job_id}"
    if format=="txt":
        return Response(content=text,media_type="text/plain; charset=utf-8",headers={"Content-Disposition":f'attachment; filename="{filename}.txt"'})
    if format=="pdf":
        stream=io.BytesIO()
        styles=getSampleStyleSheet()
        styles.add(ParagraphStyle(name="CVName",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=17,leading=21,textColor=colors.HexColor("#4C3933"),alignment=TA_CENTER,spaceAfter=8))
        styles.add(ParagraphStyle(name="CVHeading",parent=styles["Heading2"],fontName="Helvetica-Bold",fontSize=11,leading=14,textColor=colors.HexColor("#4C3933"),spaceBefore=8,spaceAfter=4))
        styles.add(ParagraphStyle(name="CVBody",parent=styles["BodyText"],fontName="Helvetica",fontSize=9,leading=12,spaceAfter=3))
        pdf=SimpleDocTemplate(stream,pagesize=A4,rightMargin=48,leftMargin=48,topMargin=42,bottomMargin=42,title=f"Currículo direcionado - {job.title}")
        story=[]
        heading_names={"resumo do currículo","experiência profissional","experiencia profissional","experiência","experiencia","formação acadêmica","formacao academica","educação","educacao","competências","competencias","habilidades","cursos","idiomas","resumo profissional","objetivo","nota de integridade"}
        for idx,line in enumerate(text.splitlines()):
            line=line.strip()
            if not line:
                story.append(Spacer(1,4)); continue
            safe=line.replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")
            if idx==0: story.append(Paragraph(safe,styles["CVName"]))
            elif line.lower().strip(" :") in heading_names: story.append(Paragraph(safe,styles["CVHeading"]))
            else: story.append(Paragraph(safe,styles["CVBody"]))
        pdf.build(story)
        return Response(content=stream.getvalue(),media_type="application/pdf",headers={"Content-Disposition":f'attachment; filename="{filename}.pdf"'})
    doc=Document()
    for line in text.splitlines():
        if line.strip() and line.strip().lower().strip(" :") in {"resumo do currículo","experiência profissional","experiencia profissional","experiência","experiencia","formação acadêmica","formacao academica","educação","educacao","competências","competencias","habilidades","cursos","idiomas","resumo profissional","objetivo","nota de integridade"}:
            doc.add_heading(line.strip(),level=2)
        elif line.strip():
            doc.add_paragraph(line.strip())
    stream=io.BytesIO(); doc.save(stream)
    return Response(content=stream.getvalue(),media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",headers={"Content-Disposition":f'attachment; filename="{filename}.docx"'})

@app.get("/api/jobs/{job_id}/analysis")
def analyze_job(job_id:int,user=Depends(require_auth)):
    with Session() as db:
        job=db.scalar(select(Job).where(Job.id==job_id,Job.user_id==user.id)); profile=db.scalar(select(Profile).where(Profile.user_id==user.id))
        if not job or not profile: raise HTTPException(404,"Vaga não encontrada")
        cv=(profile.resume_text or "").lower(); desc=(job.description or "").lower(); title=(job.title or "").lower()
        skill_terms=["marketplace","e-commerce","ecommerce","social media","instagram","tiktok","shopee","mercado livre","catalog","catálogo","ads","anúncios","atendimento","customer support","excel","canva","content","conteúdo","estoque","logística","product listing","gestão","gestao"]
        present=[term for term in skill_terms if term in cv and (term in desc or term in title)]
        missing=[term for term in skill_terms if term in desc and term not in cv]
        lang_required=[term for term in ["fluent in spanish","spanish fluency","fluency in spanish","advanced spanish","spanish required","fluent in english","english fluency","advanced english required","english required","espanhol fluente","inglês fluente","ingles fluente"] if term in desc]
        return {"job_id":job.id,"title":job.title,"company":job.company,"score":job.score,"matched_terms":present,"possible_gaps":missing,"language_warnings":lang_required,"description_available":bool(job.description),"note":"Análise heurística baseada no texto disponível; requisitos ausentes no currículo não provam falta de competência."}

@app.get("/api/jobs")
def list_jobs(user=Depends(require_auth)):
    with Session() as db:
        jobs=db.scalars(select(Job).where(Job.user_id==user.id,Job.score>0).order_by(Job.score.desc(),Job.id.desc())).all()
        return [dict(id=j.id,url=j.url,title=j.title,company=j.company,location=j.location,source=j.source,salary=j.salary,contract_type=j.contract_type,location_check=location_verification_label(j.location),score=j.score,status=j.status,created_at=j.created_at.isoformat()) for j in jobs]

@app.post("/api/search")
async def search_jobs(user=Depends(require_auth), limit: int = 50):
    limit=max(1,min(limit,100)); found=[]; errors=[]
    async with httpx.AsyncClient(timeout=20,follow_redirects=True,headers={"User-Agent":"VagasIA/1.0 job search assistant"}) as client:
        try:
            r=await client.get("https://jobicy.com/api/v2/remote-jobs",params={"count":limit,"geo":"brazil"}); r.raise_for_status()
            for j in r.json().get("jobs",[]):
                found.append({"url":j.get("url"),"title":j.get("jobTitle") or "Vaga sem título","company":j.get("companyName") or "Empresa não informada","location":j.get("jobGeo") or "Remoto","source":"Jobicy","description":j.get("jobDescription") or j.get("jobExcerpt") or "", "salary":j.get("salary") or j.get("jobSalary") or "", "contract_type":j.get("jobType") or ""})
        except Exception as e: errors.append("Jobicy: "+str(e)[:120])
        try:
            r=await client.get("https://remotive.com/api/remote-jobs",params={"limit":limit}); r.raise_for_status()
            for j in r.json().get("jobs",[]):
                loc=j.get("candidate_required_location") or ""
                if any(x in loc.lower() for x in ["worldwide","anywhere","brazil","brasil","south america","latin america","all countries","global"]):
                    found.append({"url":j.get("url"),"title":j.get("title") or "Vaga sem título","company":j.get("company_name") or "Empresa não informada","location":loc or "Remoto","source":"Remotive","description":j.get("description") or "", "salary":j.get("salary") or "", "contract_type":j.get("job_type") or ""})
        except Exception as e: errors.append("Remotive: "+str(e)[:120])
    inserted=0
    filtered_existing=0
    with Session() as db:
        p=db.scalar(select(Profile).where(Profile.user_id==user.id))
        if p is None: raise HTTPException(404,"Perfil não encontrado")
        # Re-score only jobs owned by this account.
        for saved in db.scalars(select(Job).where(Job.user_id==user.id)).all():
            new_score=score_job(saved.title,saved.location,p.target_roles,p.skills,p.resume_text,saved.description)
            if salary_is_below_minimum(saved.salary, minimum_salary_brl(p.target_salary)): new_score=0
            if new_score <= 0 and saved.score > 0:
                saved.score=0
                filtered_existing += 1
            elif new_score > 0:
                saved.score=new_score
        for j in found:
            if not j.get("url"): continue
            score=score_job(j["title"],j["location"],p.target_roles,p.skills,p.resume_text,j.get("description",""))
            if salary_is_below_minimum(j.get("salary",""), minimum_salary_brl(p.target_salary)): score=0
            old=db.scalar(select(Job).where(Job.user_id==user.id,Job.url==j["url"]))
            if score<=0:
                if old and old.score>0:
                    old.score=0
                    filtered_existing += 1
                continue
            old=db.scalar(select(Job).where(Job.user_id==user.id,Job.url==j["url"]))
            if old:
                old.title=j["title"]; old.company=j["company"]; old.location=j["location"]; old.source=j["source"]; old.description=j.get("description",""); old.salary=j.get("salary",""); old.contract_type=j.get("contract_type",""); old.score=score
            else:
                db.add(Job(user_id=user.id,url=j["url"],title=j["title"],company=j["company"],location=j["location"],source=j["source"],description=j.get("description",""),salary=j.get("salary",""),contract_type=j.get("contract_type",""),score=score,status="Encontrada")); inserted+=1
        db.commit()
        total=db.scalar(select(Job).where(Job.score>0).with_only_columns(Job.id).order_by(None).subquery().select()) if False else len(db.scalars(select(Job.id).where(Job.user_id==user.id,Job.score>0)).all())
    return {"new":inserted,"total":total,"checked":len(found),"filtered_existing":filtered_existing,"errors":errors}

@app.patch("/api/jobs/{job_id}")
def update_job(job_id:int,payload:StatusUpdate,user=Depends(require_auth)):
    allowed={"Encontrada","Prioritária","Pronta para candidatura","Candidatura em andamento","Candidatura confirmada","Entrevista","Recusada","Encerrada","Erro de envio"}
    if payload.status not in allowed: raise HTTPException(400,"Status inválido")
    with Session() as db:
        j=db.scalar(select(Job).where(Job.id==job_id,Job.user_id==user.id))
        if not j: raise HTTPException(404,"Vaga não encontrada")
        j.status=payload.status; db.commit()
    return {"ok":True}

@app.get("/api/stats")
def stats(user=Depends(require_auth)):
    with Session() as db:
        jobs=db.scalars(select(Job).where(Job.user_id==user.id)).all()
        return {"total":sum(1 for j in jobs if j.score>0),"priority":sum(1 for j in jobs if j.score>=50),"applied":sum(1 for j in jobs if j.status=="Candidatura confirmada")}

@app.get("/api/export.csv")
def export_csv(user=Depends(require_auth)):
    output=io.StringIO(); writer=csv.writer(output)
    writer.writerow(["Vaga","Empresa","Localização","Salário informado","Vínculo informado","Fonte","Link","Compatibilidade","Status","Encontrada em"])
    with Session() as db:
        for j in db.scalars(select(Job).where(Job.user_id==user.id).order_by(Job.score.desc())).all():
            writer.writerow([j.title,j.company,j.location,j.salary or "Não informado",j.contract_type or "Não informado",j.source,j.url,j.score,j.status,j.created_at.isoformat() if j.created_at else ""])
    return Response(content="\ufeff"+output.getvalue(),media_type="text/csv; charset=utf-8",headers={"Content-Disposition":"attachment; filename=vagas_ia.csv"})
